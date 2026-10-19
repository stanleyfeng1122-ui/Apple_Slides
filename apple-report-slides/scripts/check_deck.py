#!/usr/bin/env python3
"""Check a native deck against its PNGs and an exported PDF, in 1920×1080 pixels.

The PDF must come from the same deck/render operation. Text is matched to native
frames by content and proximity; table bounds use rendered cell fills/rules, so
application-driven row growth is measured rather than guessed from OOXML.
Use --before with the pre-fix PNG directory to detect new empty-band regressions.
"""
import argparse
import math
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
import re
import unicodedata

import numpy as np
from PIL import Image
import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

STAGE = (1920, 1080)
WARNING = ('Glass is a baked picture: rebuild the deck after moving it, editing '
           'what is under it, or lengthening its text.')
LOCAL_PATH = re.compile(r'/Users/|/home/|~/|\b[A-Za-z]:\\|file://', re.I)


def words(text):
    return [w for w in text.split() if any(c.isalnum() for c in w)]


def normal(text):
    return ''.join(c for c in unicodedata.normalize('NFKC', text).casefold() if c.isalnum())


def contains(a, b, slack=1):
    return (a.x0-slack <= b.x0 and a.y0-slack <= b.y0 and
            a.x1+slack >= b.x1 and a.y1+slack >= b.y1)


def intersection(a, b):
    return max(0, min(a.x1, b.x1)-max(a.x0, b.x0)), max(0, min(a.y1, b.y1)-max(a.y0, b.y0))


@dataclass
class Element:
    name: str
    box: pymupdf.Rect
    kind: str
    text: str = ''
    parent: str = ''
    lines: list = field(default_factory=list)
    tokens: set = field(default_factory=set)
    chart_layouts: list = field(default_factory=list)


def scaled(rect, sx, sy):
    return pymupdf.Rect(rect[0]*sx, rect[1]*sy, rect[2]*sx, rect[3]*sy)


def collect(slide, deck, page):
    """Resolve nested group transforms and inherited footer furniture."""
    sx, sy = STAGE[0]/deck.slide_width, STAGE[1]/deck.slide_height
    px, py = STAGE[0]/page.rect.width, STAGE[1]/page.rect.height
    paths = [scaled(d['rect'], px, py) for d in page.get_drawings()]
    elements, tables = [], []

    def walk(shapes, transform=(1, 1, 0, 0)):
        ax, ay, ox, oy = transform
        for shape in shapes:
            x, y = ox+shape.left*ax, oy+shape.top*ay
            w, h = shape.width*ax, shape.height*ay
            box = scaled((x,y,x+w,y+h),sx,sy)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                tr = shape._element.grpSpPr.xfrm
                gx = tr.ext.cx/tr.chExt.cx if tr.chExt.cx else 1
                gy = tr.ext.cy/tr.chExt.cy if tr.chExt.cy else 1
                walk(shape.shapes, (ax*gx,ay*gy,x-tr.chOff.x*ax*gx,y-tr.chOff.y*ay*gy))
                continue
            if shape.has_table:
                table = shape.table
                edges = [box.x0]
                for col in table.columns:
                    edges.append(edges[-1]+col.width*ax*sx)
                # Cell rectangles are identifiable by their exact column edges.
                rects = [r for r in paths if r.y0 >= box.y0-2 and r.y0 < 980 and
                         r.height > 4 and any(abs(r.x0-a)<2 and abs(r.x1-b)<2
                         for a,b in zip(edges,edges[1:]))]
                boundaries = [box.y0]
                for row in table.rows:
                    start = boundaries[-1]
                    candidates = [r.y1 for r in rects if abs(r.y0-start)<3]
                    boundaries.append(max(candidates) if candidates else start+row.height*ay*sy)
                actual = pymupdf.Rect(box.x0,box.y0,box.x1,boundaries[-1])
                owner = Element(shape.name,actual,'table')
                tables.append(owner)
                elements.append(owner)
                for ri,row in enumerate(table.rows):
                    for ci,cell in enumerate(row.cells):
                        rect = pymupdf.Rect(edges[ci],boundaries[ri],edges[ci+1],boundaries[ri+1])
                        elements.append(Element(f'{shape.name} R{ri+1}C{ci+1}',rect,'cell',cell.text,shape.name))
                continue
            kind = 'chart' if shape.has_chart else 'picture' if shape.shape_type==MSO_SHAPE_TYPE.PICTURE else 'shape'
            text = shape.text if shape.has_text_frame else ''
            element = Element(shape.name,box,kind,text)
            if shape.has_chart:
                element.chart_layouts = shape.chart._chartSpace.xpath('.//c:manualLayout')
            elements.append(element)

    walk(slide.shapes)
    # Layout/master placeholder text is overridden by slide placeholders. Footer
    # marks and number fields are real inherited shapes in these report layouts.
    keys = {(e.name,tuple(round(v,1) for v in e.box)) for e in elements}
    for inherited in (slide.slide_layout, slide.slide_layout.slide_master):
        for shape in inherited.shapes:
            if shape.is_placeholder or not any(label in shape.name.casefold() for label in ('footer', 'classification', 'slide number')):
                continue
            box = scaled((shape.left,shape.top,shape.left+shape.width,shape.top+shape.height),sx,sy)
            key = shape.name,tuple(round(v,1) for v in box)
            if key not in keys:
                name = {'Classification': 'footer-mark', 'Slide number': 'footer-number', 'Footer rule': 'footer-rule'}.get(shape.name, shape.name)
                text = str(deck.slides.index(slide)+1) if name == 'footer-number' else shape.text if shape.has_text_frame else ''
                elements.append(Element(name,box,'shape',text))
                keys.add(key)
    for e in elements:
        e.tokens = {normal(w) for w in e.text.split() if normal(w)}
    # Keep PDF line IDs: independent spans on a single line must remain one line.
    mapped = defaultdict(list)
    for word in page.get_text('words'):
        box = scaled(word[:4],px,py)
        token = normal(word[4])
        candidates = [e for e in elements if e.text and (token in e.tokens if token else word[4] in e.text)]
        if not candidates:
            continue
        if any(e.kind == 'chart' and contains(e.box, box) for e in elements):
            candidates = [e for e in candidates if intersection(e.box, box)[0] > 0 and intersection(e.box, box)[1] > 0]
            if not candidates:
                continue
        def score(e):
            dx = max(e.box.x0-box.x1,box.x0-e.box.x1,0)
            dy = max(e.box.y0-box.y1,box.y0-e.box.y1,0)
            return dx*4+dy+abs(e.box.x0-box.x0)*.015+abs(e.box.y0-box.y0)*.01
        owner = min(candidates,key=score)
        # Unrelated chart labels must not attach to distant matching numerals.
        if score(owner)>400:
            continue
        mapped[(id(owner),word[5],word[6])].append(box)
    for e in elements:
        for (eid,_,_),boxes in mapped.items():
            if eid==id(e):
                line=pymupdf.Rect(boxes[0])
                for b in boxes[1:]: line |= b
                e.lines.append(line)
    return elements,tables


def empty_band(path):
    """Largest flat-colour rectangle exceeding both limits, including card faces.

    Four-pixel blocks must be flat in every pixel (3/255 tolerance); this keeps
    even thin rules and small text occupied. A monotone histogram finds the
    maximum rectangle per colour in O(width × height), with a 4px resolution.
    """
    with Image.open(path) as image:
        if image.size != STAGE:
            raise ValueError(f'{path}: expected 1920×1080 PNG, got {image.size}')
        rgb=np.asarray(image.convert('RGB'))[200:980].astype(np.int16)
    blocks=rgb.reshape(195,4,480,4,3)
    low=blocks.min(axis=(1,3)); high=blocks.max(axis=(1,3))
    flat=(high-low).max(axis=2)<=3
    colour=(low//4).astype(np.int32)
    codes=(colour[:,:,0]*4096+colour[:,:,1]*64+colour[:,:,2])
    codes[~flat]=-1
    heights=np.zeros(480,dtype=np.int32)
    previous=np.full(480,-2,dtype=np.int32)
    best=None; area=0
    for y,row in enumerate(codes):
        heights=np.where((row==previous)&(row>=0),heights+1,np.where(row>=0,1,0))
        start=0
        while start<480:
            end=start+1
            while end<480 and row[end]==row[start]: end+=1
            if row[start]>=0:
                stack=[]
                for x in range(start,end+1):
                    height=int(heights[x]) if x<end else 0
                    left=x
                    while stack and stack[-1][1]>height:
                        pos,h=stack.pop(); left=pos
                        width=(x-pos)*4; depth=h*4
                        if width>560 and depth>300 and width*depth>area:
                            best=pymupdf.Rect(pos*4,200+(y+1-h)*4,x*4,200+(y+1)*4)
                            area=width*depth
                    if not stack or stack[-1][1]<height: stack.append((left,height))
            start=end
        previous=row
    return best


def exception_slides(outline):
    if not outline.is_file(): return set()
    result=set()
    for line in outline.read_text().splitlines():
        match = re.match(r'^\s*(?:S0*(\d+)|slide\s+0*(\d+))\b.*'
                         r'·\s*exception:\s*([^·\s][^·]*)\s*$', line, re.I)
        if match:
            result.add(int(match.group(1) or match.group(2)))
    return result


def graded(delta, action):
    if delta is None: return f'≤40 px class: metadata/text repair (0 px); {action}'
    delta=max(0,delta)
    grade=('nudge or compact spacing (≤40 px)' if delta<=40 else
           'shorten text or tighten the group (40–90 px)' if delta<=90 else
           'restructure the slide (90–160 px)' if delta<=160 else
           'switch layout or split (>160 px)')
    return f'{grade}; {action}'


def chart_findings(elements, page):
    """Check native plot bounds and actual PDF text, independently of the renderer."""
    px, py = STAGE[0]/page.rect.width, STAGE[1]/page.rect.height
    pdf_lines = []
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            text = ''.join(span['text'] for span in line['spans']).strip()
            if text:
                pdf_lines.append((scaled(line['bbox'], px, py), text))
    for chart in (e for e in elements if e.kind == 'chart'):
        for layout in chart.chart_layouts:
            values = {}
            for key in 'xywh':
                node = layout.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}' + key)
                if node is not None:
                    try:
                        values[key] = float(node.get('val'))
                    except (TypeError, ValueError):
                        values[key] = float('nan')
            invalid = any(not math.isfinite(v) or not 0 < v <= 1 for v in values.values())
            invalid |= all(k in values for k in 'xw') and values['x']+values['w'] > 1
            invalid |= all(k in values for k in 'yh') and values['y']+values['h'] > 1
            if invalid:
                yield ('chart-layout', chart.name, f'manualLayout {values}',
                       're-measure the plot or omit invalid manual layout', None)
        lines = [(box, text) for box, text in pdf_lines if contains(chart.box, box)]
        for i, (a, text_a) in enumerate(lines):
            for b, text_b in lines[i+1:]:
                w, h = intersection(a, b)
                if w > 1 and h > 1:
                    yield ('chart-text', chart.name,
                           f'PDF lines {text_a!r} / {text_b!r} overlap {w:.1f}×{h:.1f}px',
                           'separate chart labels or repair the plot layout', min(w, h))


def check(deck_path, render_dir, pdf_path, before=None):
    deck=Presentation(deck_path)
    registry={p.stem for p in (Path(__file__).resolve().parent.parent/'assets/layouts').glob('*.html')}
    exceptions=exception_slides(deck_path.parent/'outline.md')
    findings=[]
    def emit(number,check_name,severity,element,measured,action,delta=None):
        line=f'S{number:02d} · {check_name} · {severity} · {element} · {measured} · {graded(delta,action)}'
        findings.append((severity,line))
    pymupdf.TOOLS.set_small_glyph_heights(True)
    with pymupdf.open(pdf_path) as pdf:
        if len(pdf)!=len(deck.slides): raise ValueError('PDF page count differs from PPTX slide count')
        layouts=[]
        for n,(slide,page) in enumerate(zip(deck.slides,pdf),1):
            name=slide.name.strip(); layout=name.split(' · ',1)[0]
            layouts.append(layout)
            if layout not in registry:
                emit(n,'layout-id','FAIL','slide name',repr(name),'stamp a registered layout ID, optionally followed by · variant')
            if n>=3 and layout in registry and layouts[-3:]==[layout]*3:
                emit(n,'variety','WARN','slides',f'S{n-2:02d}–S{n:02d}: {layout}','choose a different fitting layout',161)
            elements,tables=collect(slide,deck,page)
            for check_name, element, measured, action, delta in chart_findings(elements, page):
                emit(n, check_name, 'FAIL', element, measured, action, delta)
            notes=slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ''
            # Native text retains paths split across rendered lines; PDF text
            # also covers chart labels and inherited text outside collect().
            visible_text='\n'.join(e.text for e in elements)+'\n'+page.get_text()
            for location,text in (('slide text',visible_text),('presenter notes',notes)):
                match=LOCAL_PATH.search(text)
                if match:
                    emit(n,'notes-path','FAIL',location,f'local path or file link ({match.group(0)!r})',
                         'cite sources by title plus page, section or heading; remove local paths and file links')
            if not notes or len(words(notes))>150:
                emit(n,'notes','FAIL','presenter notes',f'{len(words(notes))} words','write a non-empty script of ≤150 words')
            glass=any(e.kind == 'picture' and 'glass' in e.name.casefold() and 'rebuild' in e.name.casefold() for e in elements)
            if glass and not notes.endswith(WARNING):
                emit(n,'notes','FAIL','glass warning','missing final rebuild warning','end notes with the exact glass warning')
            marks=[e for e in elements if 'footer-mark' in e.name]
            if not any(e.text.strip() and e.lines for e in marks):
                emit(n,'mark','FAIL','footer classification','missing or not rendered','restore a visible, non-empty source classification mark')
            count=0
            for e in elements:
                if 'footer-number' in e.name: continue
                if e.kind=='cell' and not any(c.isalpha() for c in e.text): continue
                count+=len(words(e.text))
            if layout!='title-evidence' and count>90:
                emit(n,'words','WARN','visible text',f'{count} counted words >90','shorten prose; retain evidence and move context to notes')
            for e in elements:
                if not e.lines: continue
                overflow=max(max(e.box.x0-r.x0,e.box.y0-r.y0,r.x1-e.box.x1,r.y1-e.box.y1) for r in e.lines)
                footer=max(r.y1-1004 for r in e.lines) if 'footer' not in e.name else 0
                if overflow>3 or footer>0:
                    emit(n,'overflow','FAIL',e.name,f'past frame {max(0,overflow):.1f}px; past footer {max(0,footer):.1f}px',
                         'fit the rendered lines inside their frame and above y=1004',max(overflow,footer))
            pairs=set()
            for e in elements:
                for other in elements:
                    if e is other or not e.lines or e.parent==other.name or other.parent==e.name: continue
                    if contains(other.box,e.box) or contains(e.box,other.box): continue
                    if e.parent and e.parent==other.parent: continue
                    # Table cells, not the enclosing table, own text collisions.
                    if other.kind=='table': continue
                    targets=other.lines if other.text and other.lines else [other.box]
                    clashes=[intersection(a,b) for a in e.lines for b in targets]
                    clashes=[min(w,h) for w,h in clashes if w>1 and h>1]
                    pair=tuple(sorted((e.name,other.name)))
                    if clashes and pair not in pairs:
                        pairs.add(pair)
                        size=max(clashes)
                        emit(n,'overlap','FAIL',f'{e.name} / {other.name}',f'line intersection {size:.1f}px','separate these elements without creating a new empty band',size+8)
            for table in tables:
                candidates=[e for e in elements if e is not table and e.parent!=table.name and 'footer' not in e.name and
                            e.kind!='cell' and e.box.y0>=table.box.y0 and intersection(e.box,table.box)[0]>1 and
                            not contains(e.box,table.box) and not contains(table.box,e.box)]
                if not candidates: continue
                following=min(candidates,key=lambda e:e.box.y0)
                gap=following.box.y0-table.box.y1
                if gap<8:
                    emit(n,'touching','FAIL',f'{table.name} → {following.name}',f'rendered bottom {table.box.y1:.1f}px; gap {gap:.1f}px <8',
                         'leave at least 8px; prefer a body line pitch +8px for edit room',8-gap)
            band=empty_band(render_dir/f'S{n:02d}.png')
            if band:
                containers=[e for e in elements if contains(e.box,band) and e.kind=='shape']
                owner=min(containers,key=lambda e:e.box.get_area()).name if containers else 'content band'
                overcorrection = False
                if before is not None:
                    with Image.open(before/f'S{n:02d}.png') as old_image:
                        if old_image.size != STAGE:
                            raise ValueError('Pre-fix PNG must be 1920×1080')
                        old = np.asarray(old_image.convert('RGB').crop(tuple(band))).astype(np.int16)
                    overcorrection = bool((old.max(axis=(0, 1))-old.min(axis=(0, 1))).max() > 3)
                severity='INFO' if n in exceptions else 'WARN'
                emit(n,'empty-band',severity,owner,
                     f'{band.width:.0f}×{band.height:.0f}px at ({band.x0:.0f},{band.y0:.0f})'+
                     ('; over-correction: new after fix' if overcorrection else '')+
                     ('; outline exception' if n in exceptions else ''),
                     'compact the group, add supplied evidence or merge; never pad',min(band.width-560,band.height-300))
    for _,line in findings: print(line)
    print(f'Check: {len(deck.slides)} slides · '+ ' · '.join(f'{sum(s==label for s,_ in findings)} {label}' for label in ('FAIL','WARN','INFO')))
    return 1 if any(s=='FAIL' for s,_ in findings) else 0


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('deck',type=Path)
    parser.add_argument('render_dir',type=Path)
    parser.add_argument('--pdf',required=True,type=Path)
    parser.add_argument('--before',type=Path,help='Pre-fix PNGs; label new empty bands as over-corrections')
    args=parser.parse_args()
    try:
        return check(args.deck.resolve(),args.render_dir.resolve(),args.pdf.resolve(),args.before)
    except (OSError,ValueError,KeyError,pymupdf.FileDataError) as error:
        parser.exit(1,f'Check failed: {error}\n')


if __name__=='__main__':
    raise SystemExit(main())

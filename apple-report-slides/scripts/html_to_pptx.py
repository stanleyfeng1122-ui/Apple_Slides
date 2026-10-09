#!/usr/bin/env python3
"""Convert marked local HTML pages to editable PowerPoint using rendered geometry.

Requires already-installed Playwright, Google Chrome and python-pptx.
No browser download, HTTP server or external resource access is used.
"""
import sys
sys.dont_write_bytecode = True

import argparse
import base64
from io import BytesIO
from pathlib import Path
import re
import uuid

import glass

from lxml import etree
from urllib.parse import unquote, urlparse

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_MARK, XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.shapes.autoshape import Shape
from pptx.util import Pt


EXTRACT = r"""page => {
  const origin = page.getBoundingClientRect();
  const number = value => parseFloat(value) || 0;
  function box(el) {
    const r = el.getBoundingClientRect();
    return {x:r.x-origin.x, y:r.y-origin.y, w:r.width, h:r.height};
  }
  function style(el) {
    const s = getComputedStyle(el);
    const border = side => ({width:number(s[`border${side}Width`]),
      color:s[`border${side}Color`], style:s[`border${side}Style`]});
    return {font:s.fontFamily.split(',')[0].replaceAll('"','').trim(),
      size:number(s.fontSize), weight:number(s.fontWeight), color:s.color,
      italic:s.fontStyle==='italic', align:s.textAlign, vertical:s.verticalAlign,
      baseline:el.closest('sup') ? 35000*number(getComputedStyle(el.closest('sup').parentElement).fontSize)/number(s.fontSize)
        : el.closest('sub') ? -25000*number(getComputedStyle(el.closest('sub').parentElement).fontSize)/number(s.fontSize) : 0,
      leading:number(s.lineHeight)||number(s.fontSize)*1.2,
      spacing:number(s.letterSpacing), fill:s.backgroundColor, textTransform:s.textTransform,
      radius:number(s.borderTopLeftRadius),
      padding:[s.paddingTop,s.paddingRight,s.paddingBottom,s.paddingLeft].map(number),
      border:{top:border('Top'),right:border('Right'),bottom:border('Bottom'),left:border('Left')}};
  }
  // Hard breaks become paragraphs; browser soft wraps remain one paragraph.
  function textData(el, cell=false) {
    const paragraphs=[[]], visual=[];
    function append(text, s) {
      if (!text) return;
      if(s.textTransform==='uppercase') text=text.toUpperCase();
      if(s.textTransform==='lowercase') text=text.toLowerCase();
      const runs=paragraphs.at(-1);
      const key=JSON.stringify([s.font,s.size,s.weight,s.color,s.italic,s.spacing,s.baseline]);
      if(runs.at(-1)?.key===key) runs.at(-1).text+=text;
      else runs.push({text,style:s,key});
    }
    function walk(node) {
      if(node.nodeType===Node.ELEMENT_NODE) {
        const cs=getComputedStyle(node);
        if(cs.display==='none'||cs.visibility==='hidden'||node.dataset.role==='image') return;
        if(cell && node!==el && node.matches('.status-pill,[data-role="shape"]')) return;
        if(node.tagName==='BR') { paragraphs.push([]); return; }
        if(node!==el && !['inline','inline-block','contents'].includes(cs.display)) throw new Error('Block inside marked text: use <br> for a hard break');
        node.childNodes.forEach(walk); return;
      }
      if(node.nodeType!==Node.TEXT_NODE) return;
      const parent=node.parentElement, css=getComputedStyle(parent), st=style(parent);
      const hard=/pre|break-spaces/.test(css.whiteSpace);
      const parts=(hard ? node.textContent.replace(/\r\n?/g,'\n') : node.textContent.replace(/\s+/g,' ')).split('\n');
      parts.forEach((part,i)=>{if(i) paragraphs.push([]);append(part,st);});
      for(let i=0;i<node.textContent.length;i++) {
        const range=document.createRange();range.setStart(node,i);range.setEnd(node,i+1);
        const r=range.getBoundingClientRect();
        if(r.width<0.01||/\s/.test(node.textContent[i])) continue;
        if(!visual.some(v=>Math.min(v.bottom,r.bottom)-Math.max(v.top,r.top)>Math.min(v.height,r.height)*0.3)) visual.push(r);
      }
    }
    walk(el);
    while(!cell && paragraphs.length>1 && !paragraphs[0].some(r=>r.text.trim())) paragraphs.shift();
    while(paragraphs.length>1 && !paragraphs.at(-1).some(r=>r.text.trim())) paragraphs.pop();
    return {lines:paragraphs,wrap:true,visualLines:visual.length};
  }
  // One-line width of the element's content box, without wrapping.
  function naturalWidth(el, padding=0) {
    const clone=el.cloneNode(true);
    Object.assign(clone.style,{position:'absolute',width:'max-content',maxWidth:'none',
      height:'auto',whiteSpace:'nowrap',visibility:'hidden'});
    el.parentElement.append(clone);
    const width=clone.getBoundingClientRect().width-padding;
    clone.remove();
    return width;
  }
  function textGeometry(el, item) {
    const s=getComputedStyle(el), b=item.box;
    const pill=el.matches('.status-pill,.pill,.badge,.plain-callout,[data-inline="true"]');
    const numeric=/^[\s\d.,%+−–—/<>≤≥±×:$€£()]+$/.test(el.textContent);
    const glassCaption=el.parentElement.matches('.glass,[data-role="glass"]') &&
      el.matches('.g-cap,.caption,.g-lab,.g-tag');
    item.inline=!glassCaption && (pill || numeric || ['inline','inline-block'].includes(s.display));
    item.wrap=!item.inline;
    if(pill) {
      const padding=item.style.padding[1]+item.style.padding[3];
      const textWidth=naturalWidth(el,padding);
      const required=textWidth+padding+Math.max(6,textWidth*.08);
      const width=Math.max(b.w,required);
      b.x-=(width-b.w)/2;b.w=width;
      item.textWidth=textWidth;item.slack=Math.max(6,textWidth*.08);
      item.style.align='center';
    } else if(!item.inline) {
      // Explicit widths describe authored columns. Auto-width absolute text is
      // shrink-wrapped by Chrome; give it the containing block's content width.
      const authored=el.style.width || [...document.styleSheets].some(sheet=>{
        try {return [...sheet.cssRules].some(rule=>rule.selectorText && rule.style?.width &&
          el.matches(rule.selectorText));} catch {return false;}
      });
      if(!authored && s.position==='absolute') {
        const parent=el.offsetParent || el.parentElement, ps=getComputedStyle(parent), pb=box(parent);
        const right=pb.x+pb.w-number(ps.paddingRight)-number(ps.borderRightWidth);
        b.w=Math.max(1,right-b.x);
      }
      // Glass captions retain the capsule's inner width even when wrapped.
      // The baked face cannot grow: longer copy requires reauthoring/rebuilding.
      // Plain callouts keep the existing single-line width compensation.
      const container=el.parentElement;
      if(glassCaption || container.matches('.glass,[data-role="glass"],[class*="callout"]') && item.visualLines<=item.lines.length) {
        const cs=getComputedStyle(container), cb=box(container);
        const left=cb.x+number(cs.paddingLeft)+number(cs.borderLeftWidth);
        const right=cb.x+cb.w-number(cs.paddingRight)-number(cs.borderRightWidth);
        if(['right','end'].includes(s.textAlign)) {b.w=b.x+b.w-left;b.x=left;}
        else if(s.textAlign==='center') {
          const mid=b.x+b.w/2, half=Math.min(mid-left,right-mid);b.x=mid-half;b.w=2*half;
        } else b.w=right-b.x;
      }
      item.containerWidth=b.w;
    }
    if(el.matches('.footer-mark')) item.textWidth=naturalWidth(el);
  }
  const walker=document.createTreeWalker(page,NodeFilter.SHOW_TEXT);
  let node;
  while(node=walker.nextNode()) {
    const el=node.parentElement, cs=getComputedStyle(el);
    if(!node.textContent.trim()||cs.display==='none'||cs.visibility==='hidden'||!el.getClientRects().length) continue;
    if(!el.closest('[data-role]')) throw new Error(`Visible text has no data-role ancestor: ${node.textContent.trim().slice(0,100)}`);
  }
  const cards=Array.from(page.querySelectorAll('.card'));
  function identity(el) {
    const card=el.closest('.card');
    return {name:el.dataset.name||el.className||el.dataset.role,
      imageNeeded:el.classList.contains('image-needed'),
      title:el.classList.contains('headline'),slideNumber:el.classList.contains('footer-number'),
      group:card ? cards.indexOf(card) : null,
      groupName:card ? card.dataset.name||`card ${cards.indexOf(card)+1}` : null};
  }

  const elements = [];
  for (const el of page.querySelectorAll('[data-role]')) {
    const role = el.dataset.role;
    if (role==='notes' || el.closest('.optic') || el.parentElement.closest('[data-role="chart"]')) continue;
    if(el.closest('[data-role="table"]') && !['table','image'].includes(role)) continue;
    const s = getComputedStyle(el);
    if (s.display==='none' || s.visibility==='hidden') continue;
    const item = {role,box:box(el),style:style(el),...identity(el)};
    if (role==='text' || role==='shape' && el.textContent.trim()) {
      Object.assign(item,textData(el));textGeometry(el,item);
    }
    else if (role==='table') {
      item.rows=Array.from(el.rows).map(row=>({height:row.getBoundingClientRect().height,
        body:row.parentElement.tagName!=='THEAD',
        cells:Array.from(row.cells).map(cell=>{
          const cs=style(cell);
          // CSS backgrounds on section/total/highlight rows are inherited
          // visually, although getComputedStyle(td).background is transparent.
          for(let parent=cell;parent && parent!==el.parentElement;parent=parent.parentElement) {
            const fill=getComputedStyle(parent).backgroundColor;
            if(fill!=='transparent' && fill!=='rgba(0, 0, 0, 0)') {cs.fill=fill;break;}
          }
          return {box:box(cell),style:cs,...textData(cell,true),rowspan:cell.rowSpan,colspan:cell.colSpan};
        })}));
    } else if (role==='chart') {
      item.chart=JSON.parse(el.dataset.chart);
      const resolveColor=value=>{
        const probe=document.createElement('span');probe.style.color=value;el.append(probe);
        const resolved=getComputedStyle(probe).color;probe.remove();return resolved;
      };
      for(const series of item.chart.series) {
        series.color=resolveColor(series.color);
        if(series.point_colors) series.point_colors=series.point_colors.map(resolveColor);
      }
      if(item.chart.label_color) item.chart.label_color=resolveColor(item.chart.label_color);
      const d=item.chart, horizontal=['bar','stacked_bar'].includes(d.type);
      const children=[...el.children].map(e=>({box:box(e),text:e.textContent.trim(),
        fill:getComputedStyle(e).backgroundColor}));
      // Measure the actual grid spans; do not guess the native chart engine's
      // label reservation. Both old decks and new template plots are supported.
      const lines=children.filter(c=>!c.text && (horizontal ? c.box.w<=2 && c.box.h>40 : c.box.h<=2 && c.box.w>40));
      if(d.category_labels && lines.length>=2) {
        const coords=lines.map(c=>horizontal?c.box.x:c.box.y).sort((a,b)=>a-b);
        const range=d.type.startsWith('stacked_')?100:d.max-d.min;
        const span=(coords[1]-coords[0])*range/(d.major_unit||range/4);
        item.plotBox=horizontal ? {x:coords[0],y:lines[0].box.y,w:span,h:lines[0].box.h}
          : {x:lines[0].box.x,y:coords.at(-1)-span,w:lines[0].box.w,h:span};
      } else if(el.dataset.plotRect) {
        const [x,y,w,h]=el.dataset.plotRect.split(',').map(Number);
        item.plotBox={x:item.box.x+x,y:item.box.y+y,w,h};
      } else {
        item.plotBox={...item.box};
      }
      const fills=d.series.flatMap(series=>[series.color,...(series.point_colors||[])]);
      const bars=children.filter(c=>!c.text && c.box.w>2 && c.box.h>2 && fills.includes(c.fill));
      if(d.type!=='line' && bars.length) {
        const thickness=bars[0].box[horizontal?'h':'w'];
        const pitch=item.plotBox[horizontal?'h':'w']/d.categories.length;
        const count=d.type.startsWith('stacked_')?1:d.series.length;
        item.gapWidth=Math.max(0,Math.min(500,Math.round(100*(pitch/thickness-count))));
      }
      // Six-digit values use thousands on the axis and labels, in Chrome and PowerPoint alike.
      if(!d.type.startsWith('stacked_') && d.series.some(series=>series.values.some(v=>Math.abs(v)>=100000)))
        item.valueFormat='#,##0,"K"';
    }
    else if (role==='image') {
      item.src=el.currentSrc||el.src;
      item.fit=el.style.objectFit || (s.objectFit==='fill' ? 'contain' : s.objectFit);
      item.natural=[el.naturalWidth,el.naturalHeight];
    } else if (role==='glass') {}
    else if (role!=='shape') throw new Error(`Unsupported role: ${role}`);
    elements.push(item);
    if(role==='table') for(const pill of el.querySelectorAll('.status-pill,[data-role="shape"]')) {
      const p={role:'shape',box:box(pill),style:style(pill),...identity(pill),...textData(pill)};
      textGeometry(pill,p);elements.push(p);
    }
  }
  // Reserve the emitted body's line pitch plus 8px for a cell gaining a line.
  // Follow column geometry across groups, moving whole following cards together.
  const isFooter=e=>String(e.name).includes('footer');
  for(const table of elements.filter(e=>e.role==='table').sort((a,b)=>a.box.y-b.box.y)) {
    // PowerPoint sizes the table from its row heights, not Chrome's table box.
    const bottom=table.box.y+Math.max(table.box.h,table.rows.reduce((sum,row)=>sum+row.height,0));
    const bodyRows=table.rows.filter(row=>row.body);
    const reserve=Math.max(28,...(bodyRows.length?bodyRows:table.rows)
      .flatMap(row=>row.cells.map(cell=>cell.style.leading+8)));
    const candidates=elements.filter(e=>e!==table && !isFooter(e) &&
      e.box.y>=table.box.y+table.box.h-1 && e.box.y<980 &&
      e.box.x<table.box.x+table.box.w && e.box.x+e.box.w>table.box.x);
    const following=new Set();
    for(const e of candidates) {
      if(e.group===null || e.group===table.group) following.add(e);
      else {
        const members=elements.filter(member=>member.group===e.group && !isFooter(member));
        // A card beside or enclosing the table is not a following card.
        if(Math.min(...members.map(member=>member.box.y))>=table.box.y+table.box.h-1)
          members.forEach(member=>following.add(member));
      }
    }
    if(!following.size) continue;
    // Text frames are emitted with add_text's leading inset; measure from there.
    const emittedTop=e=>e.box.y+(e.role==='text' ? 2-(e.style.leading-e.style.size)/2 : 0);
    const top=Math.min(...Array.from(following,emittedTop)),shift=Math.max(0,bottom+reserve-top);
    if(shift && Array.from(following).some(e=>e.role==='glass'))
      throw new Error('Leave the table body line pitch plus 8px below tables before positioning glass');
    if(shift) for(const e of following) {
      e.box.y+=shift;
      if(e.plotBox) e.plotBox.y+=shift;
    }
  }
  return {id:page.id,theme:page.dataset.theme||'report',background:getComputedStyle(page).backgroundColor,
    notes:page.querySelector('[data-role="notes"]')?.textContent||'',elements,
    tokens:Object.fromEntries(['accent','section','section-text',...Array.from(getComputedStyle(page))
      .filter(key=>/^--identity-\d+$/.test(key)).sort((a,b)=>parseInt(a.split('-').at(-1))-parseInt(b.split('-').at(-1)))
      .map(key=>key.slice(2))].map(key=>{
      const probe=document.createElement('span');probe.style.color=`var(--${key},${key==='accent'?'#0071E3':'#3C3C3E'})`;page.append(probe);
      const value=getComputedStyle(probe).color;probe.remove();return [key,value];
    }))};
}"""


def pt(px):
    """The fixed 1920 px stage maps to a 960 pt slide."""
    return Pt(float(px) / 2)


def color(value):
    if not value or value == 'transparent':
        return None
    if value.startswith('#'):
        return RGBColor.from_string(value[1:].upper())
    numbers = [float(n) for n in re.findall(r'[\d.]+', value)]
    if len(numbers) == 4 and numbers[3] == 0:
        return None
    return RGBColor(*(round(n) for n in numbers[:3]))


def paint(fill, value):
    rgb = color(value)
    if rgb is None:
        fill.background()
    else:
        fill.solid()
        fill.fore_color.rgb = rgb


def font(run, style):
    run.font.name = style['font']
    run.font.size = pt(style['size'])
    run.font.bold = style['weight'] >= 600
    run.font.italic = style['italic']
    run.font.color.rgb = color(style['color']) or RGBColor(29, 29, 31)
    if style['spacing']:
        run._r.get_or_add_rPr().set('spc', str(round(style['spacing'] * 50)))
    if style.get('baseline'):
        # Native superscript rendering shrinks glyphs to roughly 65%. The
        # browser already supplied the smaller marker size, so avoid applying
        # that reduction twice. Preserve its rise while restoring glyph size.
        run.font.size = pt(style['size'] / 0.65)
        run._r.get_or_add_rPr().set('baseline', str(round(style['baseline'] * 0.65)))


def populate(frame, lines, style, wrap=False):
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = wrap
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, runs in enumerate(lines or [[]]):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        # Empty cells still need an explicit font; the theme's large default
        # would otherwise increase their row height in native renderers.
        paragraph.font.name = style['font']
        paragraph.font.size = pt(style['size'])
        paragraph.alignment = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT,
                               'end': PP_ALIGN.RIGHT}.get(style['align'], PP_ALIGN.LEFT)
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = pt(style['leading'])
        for data in runs:
            run = paragraph.add_run()
            run.text = data['text']
            font(run, data['style'])


def add_text(slide, item):
    b, s = item['box'], item['style']
    # Compensate for native paragraph leading, using the SF Pro comparison.
    leading_inset = -(s['leading'] - s['size']) / 2 + 2
    if item.get('title') and hasattr(slide.shapes, 'title') and slide.shapes.title is not None:
        shape = slide.shapes.title
        shape.left, shape.top, shape.width, shape.height = pt(b['x']), pt(b['y'] + leading_inset), pt(b['w']), pt(b['h'])
    else:
        shape = slide.shapes.add_textbox(pt(b['x']), pt(b['y'] + leading_inset), pt(b['w']), pt(b['h']))
    populate(shape.text_frame, item['lines'], s, item.get('wrap', False))
    if item.get('slideNumber'):
        paragraph = shape.text_frame.paragraphs[0]._p
        for run in list(paragraph):
            if run.tag.endswith('}r'):
                field = OxmlElement('a:fld')
                field.set('id', '{' + str(uuid.uuid4()).upper() + '}')
                field.set('type', 'slidenum')
                for node in list(run):
                    field.append(node)
                paragraph.replace(run, field)
    return shape


def add_shape(slide, item):
    b, s = item['box'], item['style']
    if 'textWidth' in item and item.get('lines'):
        # A rounded preset shrinks its internal text rectangle in LibreOffice.
        # Keep the authored padding/slack in a separate native text frame.
        group = slide.shapes.add_group_shape()
        background = add_shape(group, {**item, 'lines': []})
        background.name = str(item['name']) + ' background'
        shape = group.shapes.add_textbox(pt(b['x']), pt(b['y']), pt(b['w']), pt(b['h']))
        shape.name = str(item['name']) + ' text'
        populate(shape.text_frame, item['lines'], s, item.get('wrap', False))
        top, right, bottom, left = s['padding']
        shape.text_frame.margin_left, shape.text_frame.margin_right = pt(left), pt(right)
        shape.text_frame.margin_top = pt(max(0, top-(s['leading']-s['size'])/2))
        shape.text_frame.margin_bottom = pt(bottom)
        if str(item.get('name', '')).startswith('callout') or 'plain-callout' in str(item.get('name', '')):
            shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
        return group
    borders = [(side, border) for side, border in s['border'].items()
               if border['width'] and border['style'] != 'none']
    if color(s['fill']) is None and len(borders) == 1 and not item.get('lines'):
        side, border = borders[0]
        coords = {'top': (0, 0, b['w'], 0), 'bottom': (0, b['h'], b['w'], b['h']),
                  'left': (0, 0, 0, b['h']), 'right': (b['w'], 0, b['w'], b['h'])}[side]
        x1, y1, x2, y2 = coords
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
            pt(b['x'] + x1), pt(b['y'] + y1), pt(b['x'] + x2), pt(b['y'] + y2))
        line.shadow.inherit = False
        for inherited in line._element.xpath('./p:style'):
            line._element.remove(inherited)
        line.line.color.rgb = color(border['color'])
        line.line.width = pt(border['width'])
        line.line.dash_style = {'dotted': MSO_LINE_DASH_STYLE.ROUND_DOT,
                                'dashed': MSO_LINE_DASH_STYLE.DASH}.get(border['style'], MSO_LINE_DASH_STYLE.SOLID)
        return line
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if s['radius'] else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, pt(b['x']), pt(b['y']), pt(b['w']), pt(b['h']))
    shape.shadow.inherit = False
    # LibreOffice still applies the theme effectRef despite an empty effectLst.
    # Every appearance property is explicit, so remove inherited shape styling.
    for inherited in shape._element.xpath('./p:style'):
        shape._element.remove(inherited)
    if s['radius']:
        shape.adjustments[0] = min(0.5, s['radius'] / min(b['w'], b['h']))
    paint(shape.fill, s['fill'])
    border = s['border']['top']
    if border['width'] and border['style'] != 'none':
        shape.line.color.rgb = color(border['color'])
        shape.line.width = pt(border['width'])
        shape.line.dash_style = {'dotted': MSO_LINE_DASH_STYLE.ROUND_DOT,
                                'dashed': MSO_LINE_DASH_STYLE.DASH}.get(border['style'], MSO_LINE_DASH_STYLE.SOLID)
    else:
        shape.line.fill.background()
    if item.get('lines'):
        populate(shape.text_frame, item['lines'], s, item.get('wrap', False))
        top, right, bottom, left = s['padding']
        shape.text_frame.margin_top = pt(max(0, top - (s['leading'] - s['size']) / 2))
        shape.text_frame.margin_right = pt(right)
        shape.text_frame.margin_bottom = pt(bottom)
        shape.text_frame.margin_left = pt(left)
        if item.get('imageNeeded') or item.get('name', '').startswith('callout') or 'plain-callout' in str(item.get('name', '')):
            shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
    return shape


def cell_border(cell, side, border):
    props = cell._tc.get_or_add_tcPr()
    name = {'top': 'lnT', 'right': 'lnR', 'bottom': 'lnB', 'left': 'lnL'}[side]
    for old in list(props):
        if old.tag.endswith('}' + name):
            props.remove(old)
    line = OxmlElement('a:' + name)
    line.set('w', str(int(pt(border['width']))))
    if border['width'] and border['style'] != 'none':
        fill = OxmlElement('a:solidFill')
        rgb = OxmlElement('a:srgbClr')
        rgb.set('val', str(color(border['color'])))
        fill.append(rgb)
        line.append(fill)
        dash = OxmlElement('a:prstDash')
        dash.set('val', {'dotted': 'sysDot', 'dashed': 'dash'}.get(border['style'], 'solid'))
        line.append(dash)
    else:
        line.append(OxmlElement('a:noFill'))
    # DrawingML requires border elements before the cell fill. PowerPoint can
    # ignore trailing border overrides and fall back to its white theme lines.
    order = ['lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr']
    rank = order.index(name)
    for index, existing in enumerate(props):
        tag = existing.tag.rsplit('}', 1)[-1]
        if tag not in order or order.index(tag) > rank:
            props.insert(index, line)
            break
    else:
        props.append(line)


def add_table(slide, item):
    rows, b = item['rows'], item['box']
    cols = len(rows[0]['cells'])
    if any(len(row['cells']) != cols or any(c['rowspan'] != 1 or c['colspan'] != 1 for c in row['cells']) for row in rows):
        raise ValueError('Only rectangular, unmerged HTML tables are supported')
    owner = slide if hasattr(slide.shapes, 'add_table') else slide.part.slide
    shape = owner.shapes.add_table(len(rows), cols, pt(b['x']), pt(b['y']), pt(b['w']), pt(b['h']))
    if owner is not slide:
        slide.shapes._spTree.insert_element_before(shape._element, 'p:extLst')
        slide.shapes._recalculate_extents()
    table = shape.table
    table.first_row = True
    table.first_col = table.last_row = table.last_col = False
    table.horz_banding = table.vert_banding = False
    for j, c in enumerate(rows[0]['cells']):
        table.columns[j].width = pt(c['box']['w'])
    for i, row in enumerate(rows):
        table.rows[i].height = pt(row['height'])
        for j, source in enumerate(row['cells']):
            cell, s = table.cell(i, j), source['style']
            paint(cell.fill, s['fill'] if color(s['fill']) else item['style']['fill'])
            populate(cell.text_frame, source['lines'], s, True)
            top, right, bottom, left = s['padding']
            cell.margin_top = pt(max(0, top - (s['leading'] - s['size']) / 2))
            cell.margin_right, cell.margin_bottom, cell.margin_left = pt(right), pt(bottom), pt(left)
            cell.vertical_anchor = {'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}.get(s.get('vertical'), MSO_ANCHOR.TOP)
            for side, border in s['border'].items():
                cell_border(cell, side, border)
    return shape


def child(parent, name, **attrs):
    el = OxmlElement(name)
    for key, value in attrs.items():
        el.set(key, str(value))
    parent.append(el)
    return el


def add_chart(slide, item):
    data, b = item['chart'], dict(item['box'])
    # Explicit optional compensation for chart-engine insets observed in the
    # local render; the HTML box remains the intended inner plot rectangle.
    left, top, right, bottom = data.get('plot_insets', [0, 0, 0, 0])
    b.update(x=b['x']-left, y=b['y']-top,
             w=b['w']+left+right, h=b['h']+top+bottom)
    types = {'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'line': XL_CHART_TYPE.LINE,
             'column': XL_CHART_TYPE.COLUMN_CLUSTERED,
             'stacked_bar': XL_CHART_TYPE.BAR_STACKED_100,
             'stacked_column': XL_CHART_TYPE.COLUMN_STACKED_100}
    if data['type'] not in types:
        raise ValueError(f"Unsupported chart type: {data['type']}")
    series_data = CategoryChartData()
    series_data.categories = data['categories']
    for series in data['series']:
        if len(series['values']) != len(data['categories']):
            raise ValueError('Chart category and value counts differ')
        series_data.add_series(series['name'], series['values'])
    shape = slide.shapes.add_chart(types[data['type']], pt(b['x']), pt(b['y']), pt(b['w']), pt(b['h']), series_data)
    chart = shape.chart
    stacked = data['type'].startswith('stacked_')
    horizontal = data['type'] in ('bar', 'stacked_bar')
    chart.has_legend = chart.has_title = False
    for axis in [chart.category_axis, chart.value_axis]:
        axis.visible = False
        axis.has_major_gridlines = False
        axis.has_minor_gridlines = False
        axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.fill.background()
    chart.value_axis.minimum_scale = 0 if stacked or data.get('category_labels') else data['min']
    if stacked:
        chart.value_axis.maximum_scale = 1
        # Authored chart ticks use percentage points (0–100); OOXML's
        # 100% stacked value axis uses fractions (0–1).
        chart.value_axis.major_unit = data.get('major_unit', 25) / 100
        chart.value_axis.tick_labels.number_format = '0%'
        chart.value_axis.tick_labels.number_format_is_linked = False
    elif 'axis_max' in data:
        largest = max(value for series in data['series'] for value in series['values'] if value is not None)
        chart.value_axis.maximum_scale = max(data['axis_max'], largest * 1.15)
    elif data.get('category_labels'):
        chart.value_axis.maximum_scale = None
    else:
        chart.value_axis.maximum_scale = data['max']
    if data.get('category_labels'):
        chart.category_axis.visible = True
        chart.value_axis.visible = True
        chart.value_axis.has_major_gridlines = True
        axis_format = item.get('valueFormat', data.get('number_format', '0'))
        if not stacked and re.fullmatch(r'[#0,]+(?:\.[0#]+)?%?', axis_format) and not axis_format.endswith(','):
            # Automatic native tick spacing is not the HTML view's range / 4.
            # General keeps automatic numeric ticks precise without trailing .0.
            if 'major_unit' not in data and '%' not in axis_format:
                axis_format = 'General'
            elif 'major_unit' in data:
                tick_scale = 100 if '%' in axis_format else 1
                ticks = (chart.value_axis.minimum_scale, data['major_unit'])
                if all(abs(v * tick_scale - round(v * tick_scale)) < 1e-9 for v in ticks):
                    axis_format = re.sub(r'\.[0#]+', '', axis_format)
        chart.value_axis.tick_labels.number_format = '0%' if stacked else axis_format
        chart.value_axis.tick_labels.number_format_is_linked = False
        if 'major_unit' in data and not stacked:
            chart.value_axis.major_unit = data['major_unit']
        chart.value_axis.major_gridlines.format.line.color.rgb = color('#D2D2D7')
        chart.value_axis.major_gridlines.format.line.width = pt(1)
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.name = 'SF Pro Text'
            axis.tick_labels.font.size = pt(data.get('label_size', 20))
            axis.tick_labels.font.color.rgb = color('#3C3C3E')
    if horizontal:
        chart.category_axis.reverse_order = True
        chart.plots[0].gap_width = item.get('gapWidth', int(data.get('gap', 100)))
    if data['type'] != 'line':
        chart.plots[0].gap_width = item.get('gapWidth', int(data.get('gap', 100)))
        if stacked:
            chart.plots[0].overlap = 100
    for series, source in zip(chart.series, data['series']):
        if data['type'] != 'line':
            paint(series.format.fill, source['color'])
            series.format.line.fill.background()
        else:
            series.format.line.color.rgb = color(source['color'])
            series.format.line.width = pt(data.get('line_width', 3))
        for point, value in zip(series.points, source.get('point_colors', [])):
            paint(point.format.fill, value)
            point.format.line.fill.background()
    if data.get('data_labels'):
        plot_labels = chart.plots[0]
        plot_labels.has_data_labels = True
        labels = plot_labels.data_labels
        labels.position = XL_DATA_LABEL_POSITION.CENTER if stacked else XL_DATA_LABEL_POSITION.OUTSIDE_END
        labels.show_value = True
        labels.show_category_name = labels.show_series_name = labels.show_legend_key = False
        labels.number_format = item.get('valueFormat', data.get('number_format', '0.0'))
        labels.font.name = item['style']['font']
        labels.font.size = pt(max(14, data.get('label_size', 20)))
        labels.font.bold = True
        labels.font.color.rgb = color(data.get('label_color', '#3C3C3E'))
    # The inner plot must match Chrome even when native axes are visible.
    plot = chart._chartSpace.chart.plotArea
    layout = plot.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}layout')
    if layout is None:
        layout = OxmlElement('c:layout')
        plot.insert(0, layout)
    manual = child(layout, 'c:manualLayout')
    child(manual, 'c:layoutTarget', val='inner')
    for key in ('xMode', 'yMode'):
        child(manual, 'c:' + key, val='edge')
    for key in ('wMode', 'hMode'):
        child(manual, 'c:' + key, val='factor')
    inner = item.get('plotBox', item['box']) if data.get('category_labels') else item['box']
    for key, value in [('x', (inner['x']-b['x'])/b['w']), ('y', (inner['y']-b['y'])/b['h']),
                       ('w', inner['w']/b['w']), ('h', inner['h']/b['h'])]:
        child(manual, 'c:' + key, val=value)
    for root in [chart._chartSpace, plot]:
        sppr = child(root, 'c:spPr')
        child(sppr, 'a:noFill')
        child(child(sppr, 'a:ln'), 'a:noFill')
    return shape


def add_image(slide, item):
    src, b = item['src'], dict(item['box'])
    parsed = urlparse(src)
    if parsed.scheme == 'data':
        header, payload = src.split(',', 1)
        if ';base64' not in header:
            raise ValueError('Image data URLs must be base64 encoded')
        source = BytesIO(base64.b64decode(payload))
    elif parsed.scheme == 'file' and parsed.netloc in ('', 'localhost'):
        source = str(Path(unquote(parsed.path)).resolve(strict=True))
    else:
        raise ValueError('Images must be local files or embedded base64 data')
    if item.get('fit', 'contain') == 'contain':
        nw, nh = item['natural']
        if not nw or not nh:
            raise ValueError(f'Image failed to load: {src}')
        factor = min(b['w'] / nw, b['h'] / nh)
        b['x'] += (b['w'] - nw * factor) / 2
        b['y'] += (b['h'] - nh * factor) / 2
        b['w'], b['h'] = nw * factor, nh * factor
    elif item['fit'] not in ('fill', 'contain'):
        raise ValueError('Images support object-fit: fill or contain, centered')
    shape = slide.shapes.add_picture(source, pt(b['x']), pt(b['y']), pt(b['w']), pt(b['h']))
    if item.get('glass'):
        glass.finish_picture(shape, item)
    elif item['style']['radius']:
        glass.finish_picture(shape, dict(box=b,glass=dict(radius=item['style']['radius'],shadow='none')))
    return shape


def extract(path, chrome, screenshots=None):
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=str(chrome), headless=True,
            args=['--disable-background-networking', '--disable-component-update', '--no-first-run'])
        try:
            def open_page(scale):
                context = browser.new_context(viewport={'width': 1920, 'height': 1080}, device_scale_factor=scale,
                                              offline=True, service_workers='block')
                context.route('**/*', lambda route: route.abort() if urlparse(route.request.url).scheme in ('http', 'https', 'ws', 'wss') else route.continue_())
                context.add_init_script("window.GLASS_DEFER_BUILD = true")
                page = context.new_page()
                page.goto(path.as_uri(), wait_until='load')
                page.evaluate('document.fonts.ready')
                glass.prepare(page, report=scale == 1)
                return context, page

            measurement, page = open_page(1)
            ids = page.locator('.page').evaluate_all('(pages)=>pages.map(p=>p.id)')
            if not ids or len(ids) != len(set(ids)) or any(not re.fullmatch(r'[A-Za-z0-9_-]+', pid) for pid in ids):
                raise ValueError('Each page needs a unique filename-safe id')
            selections = {}
            for pid in ids:
                page.evaluate('(id)=>showPage(id)', pid)
                selections[pid] = glass.calibrate(page, page.locator(f'[id="{pid}"]'))
            measurement.close()
            capture, page = open_page(2)
            result = []
            for pid in ids:
                page.evaluate('(id)=>showPage(id)', pid)
                current = page.locator(f'[id="{pid}"]')
                glass.apply_selection(page, current, selections[pid])
                data = current.evaluate(EXTRACT)
                if screenshots:
                    screenshots.mkdir(parents=True, exist_ok=True)
                    glass.screenshot(page).save(screenshots / f'{pid}.png')
                glass.bake(page, current, data)
                result.append(data)
            capture.close()
        finally:
            browser.close()
    return result


def configure_theme(deck, tokens):
    old_w, old_h = deck.slide_width, deck.slide_height
    deck.slide_width, deck.slide_height = Pt(960), Pt(540)
    deck._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz').attrib.pop('type', None)
    for part in [deck.slide_master, *deck.slide_layouts]:
        # Scale explicit placeholder geometry; inherited geometry scales on master.
        for xfrm in part._element.xpath('.//a:xfrm'):
            for node in xfrm:
                for attr in ('x', 'cx'):
                    if attr in node.attrib:
                        node.set(attr, str(round(int(node.get(attr))*deck.slide_width/old_w)))
                for attr in ('y', 'cy'):
                    if attr in node.attrib:
                        node.set(attr, str(round(int(node.get(attr))*deck.slide_height/old_h)))
    paint(deck.slide_master.background.fill, '#F3F2F5')
    for placeholder in deck.slide_master.placeholders:
        if placeholder._element.xpath('./p:nvSpPr/p:nvPr/p:ph[@type="title"]'):
            placeholder.left, placeholder.top, placeholder.width, placeholder.height = [pt(v) for v in (90,107,1740,82)]
            placeholder.text_frame.margin_left = placeholder.text_frame.margin_right = 0
            placeholder.text_frame.margin_top = placeholder.text_frame.margin_bottom = 0
    # Defaults for new slides and new text, including inherited master levels.
    for root in (deck._element, deck.slide_master._element):
        for props in root.xpath('.//a:defRPr'):
            props.set('sz', '1100')
            for old in list(props):
                if old.tag.rsplit('}', 1)[-1] in ('solidFill', 'latin', 'ea', 'cs'):
                    props.remove(old)
            child(child(props, 'a:solidFill'), 'a:schemeClr', val='tx1')
            child(props, 'a:latin', typeface='SF Pro Text')
            child(props, 'a:ea', typeface='SF Pro Text')
            child(props, 'a:cs', typeface='SF Pro Text')
    for props in deck.slide_master._element.xpath('./p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr'):
        props.set('sz', '3400')
        props.set('b', '1')
        props.getparent().set('algn', 'l')
        for face in props.xpath('./a:latin | ./a:ea | ./a:cs'):
            face.set('typeface', 'SF Pro Display')
    for part in deck.part.package.iter_parts():
        if str(part.partname).startswith('/ppt/theme/theme'):
            root = etree.fromstring(part.blob)
            ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for key, face in [('majorFont','SF Pro Display'),('minorFont','SF Pro Text')]:
                for node in root.xpath(f'//a:{key}/a:latin | //a:{key}/a:ea | //a:{key}/a:cs', namespaces=ns):
                    node.set('typeface', face)
            accent = str(color(tokens.get('accent')) or RGBColor.from_string('0071E3'))
            section = str(color(tokens.get('section')) or RGBColor.from_string('3C3C3E'))
            identities = [str(color(value)) for key, value in tokens.items()
                          if re.fullmatch(r'identity-\d+', key) and color(value)]
            deck_colours = list(dict.fromkeys([accent, *identities, section]))
            theme_accents = deck_colours[:6]
            theme_accents += [section] * (6 - len(theme_accents))
            palette = {'dk1':'3C3C3E','lt1':'F3F2F5','dk2':'3C3C3E','lt2':'FFFFFF',
                       **{f'accent{i}': value for i, value in enumerate(theme_accents, 1)},
                       'hlink':accent, 'folHlink':'6E6E73'}
            for key, value in palette.items():
                node=root.find(f'.//a:clrScheme/a:{key}', ns)
                for old in list(node): node.remove(old)
                child(node, 'a:srgbClr', val=value)
            # Keep additional source identities available without replacing the
            # six standard theme slots when a deck has more than five categories.
            extra_colours = deck_colours[6:]
            if extra_colours:
                custom = root.find('a:custClrLst', ns)
                if custom is None:
                    custom = OxmlElement('a:custClrLst')
                    extension = root.find('a:extLst', ns)
                    if extension is None:
                        root.append(custom)
                    else:
                        root.insert(root.index(extension), custom)
                for index, value in enumerate(extra_colours, 1):
                    custom_colour = child(custom, 'a:custClr')
                    custom_colour.set('name', f'Identity {index + 5}')
                    child(custom_colour, 'a:srgbClr', val=value)
            part._blob=etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def report_layouts(deck, pages):
    """Layouts for hand-added slides, using the same native typography and footer."""
    themes = {'report': ('Report content', '#F3F2F5', '#3C3C3E', '#6E6E73', '#D2D2D7'),
              'black': ('Report content — Black', '#000000', '#F5F5F7', '#A1A1A6', '#3A3A3C'),
              'white': ('Report content — White', '#FFFFFF', '#1D1D1F', '#6E6E73', '#D2D2D7')}
    result = {}
    def find(page, name):
        return next((e for e in page['elements'] if name in str(e['name'])), None)

    def classification(page):
        mark = find(page, 'footer-mark')
        return ''.join(r['text'] for line in mark['lines'] for r in line) if mark else ''

    def emitted(item):
        # Same geometry add_text gives the slide's own text frame.
        b, st = item['box'], item['style']
        return (b['x'], b['y'] - (st['leading'] - st['size']) / 2 + 2, b['w'], b['h'])

    variants = [(theme, classification(next((p for p in pages if p.get('theme', 'report') == theme), pages[0])))
                for theme in themes]
    for page in pages:
        key = (page.get('theme', 'report'), classification(page))
        if key not in variants:
            variants.append(key)
    if len(variants) > len(deck.slide_layouts):
        raise ValueError('Too many theme/classification combinations for the report layouts')
    for index, (theme, mark_text) in enumerate(variants):
        layout = deck.slide_layouts[index]
        name, paper, ink, muted, rule = themes[theme]
        if index >= 3:
            name += ' · ' + mark_text
        layout.name = name
        layout._element.set('type', 'obj')
        for shape in list(layout.shapes):
            layout.shapes._spTree.remove(shape._element)
        paint(layout.background.fill, paper)
        # Native title defaults follow the layout's text/background colour map.
        mapping = layout._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}clrMapOvr')
        if mapping is None:
            mapping = child(layout._element, 'p:clrMapOvr')
        for old in list(mapping):
            mapping.remove(old)
        colours = dict(deck.slide_master._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap').attrib)
        if theme == 'black':
            colours.update(bg1='dk1', tx1='lt1', bg2='dk2', tx2='lt2')
        child(mapping, 'a:overrideClrMapping', **colours)

        def textbox(label, geometry, text, size, face, fill, placeholder=None, align='left',
                    leading=None, spacing=0, wrap=True):
            element = CT_Shape.new_textbox_sp(len(layout.shapes)+2, label, *[pt(v) for v in geometry])
            layout.shapes._spTree.insert_element_before(element, 'p:extLst')
            if placeholder:
                kind, idx = placeholder
                child(element.nvSpPr.nvPr, 'p:ph', type=kind, idx=idx)
                element.nvSpPr.cNvSpPr.attrib.pop('txBox', None)
                child(element.nvSpPr.cNvSpPr, 'a:spLocks', noGrp=1)
            shape = Shape(element, layout.shapes)
            style = dict(font=face, size=size, weight=700 if placeholder and placeholder[0]=='title' else 400,
                         color=fill, italic=False, spacing=spacing, align=align,
                         leading=leading or (28 if size == 20 else 24))
            populate(shape.text_frame, [[dict(text=text, style=style)]], style, wrap)
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.color.rgb = color(fill)
                paragraph.font.bold = bool(placeholder and placeholder[0] == 'title')
                paragraph._p.get_or_add_pPr().set('marL', '0')
                paragraph._p.get_or_add_pPr().set('indent', '0')
            # Placeholder formatting must also be inherited by newly typed text.
            levels = element.txBody.find('{http://schemas.openxmlformats.org/drawingml/2006/main}lstStyle')
            for level in range(1, 10):
                props = child(levels, f'a:lvl{level}pPr', marL=0, indent=0, algn='l')
                child(props, 'a:buNone')
                defaults = child(props, 'a:defRPr', sz=round(size*50), b=1 if size>=64 else 0)
                if spacing:
                    defaults.set('spc', str(round(spacing * 50)))
                child(child(defaults, 'a:solidFill'), 'a:srgbClr', val=fill.lstrip('#'))
                child(defaults, 'a:latin', typeface=face)
            return shape

        # Geometry comes from this variant's own slides; template values otherwise.
        themed = [p for p in pages if p.get('theme', 'report') == theme]
        marked = [p for p in pages if classification(p) == mark_text]
        def measured(name, source):
            return next((e for e in (find(p, name) for p in source) if e), None)
        # The theme's most common headline box: title and odd slides don't set it.
        heads = [h for h in (find(p, 'headline') for p in themed) if h]
        shapes = [(emitted(h), h['style']['size']) for h in heads]
        headline = max(heads, key=lambda h: shapes.count((emitted(h), h['style']['size'])), default=None)
        if headline:
            title_box, title_size = emitted(headline), headline['style']['size']
            title_leading, title_spacing = headline['style']['leading'], headline['style']['spacing']
        elif theme == 'report':
            title_box, title_size, title_leading, title_spacing = (90, 107, 1740, 82), 68, 78, 0
        else:
            title_box, title_size, title_leading, title_spacing = (90, 100, 1740, 78), 64, 76, -.5
        textbox('Report title', title_box, 'Report headline', title_size,
                'SF Pro Display', ink, ('title', 0), leading=title_leading, spacing=title_spacing)
        textbox('Report body', (90, 254, 1740, 706), 'Add evidence', 20,
                'SF Pro Text', ink, ('body', 1))
        rule_item, number_item = measured('footer-rule', themed or pages), measured('footer-number', themed or pages)
        mark = measured('footer-mark', marked)
        rule_top = rule_item['box']['y'] if rule_item else 1004
        footer = textbox('Footer rule', (90, rule_top, 1740, 1), '', 18, 'SF Pro Text', muted)
        paint(footer.fill, rule)
        # The mark is one line, right-aligned to its measured right edge.
        if mark:
            x, y, w, _ = emitted(mark)
            width = mark['textWidth'] + max(6, mark['textWidth'] * .08)
            mark_box = (x + w - width, y, width, 26)
        else:
            mark_box = (1450, 1021, 290, 26)
        textbox('Classification', mark_box, mark_text, 18,
                'SF Pro Text', muted, align='right', wrap=False)
        number_box = emitted(number_item) if number_item else (1760, 1021, 70, 26)
        number = textbox('Slide number', number_box, '1', 18,
                         'SF Pro Text', muted, align='right', wrap=False)
        paragraph = number.text_frame.paragraphs[0]._p
        run = paragraph.find('{http://schemas.openxmlformats.org/drawingml/2006/main}r')
        field = OxmlElement('a:fld')
        field.set('id', '{' + str(uuid.uuid4()).upper() + '}')
        field.set('type', 'slidenum')
        for node in list(run):
            field.append(node)
        paragraph.replace(run, field)
        result[(theme, mark_text)] = layout
    return result


def convert(pages, output):
    deck = Presentation()
    configure_theme(deck, pages[0].get('tokens', {}))
    layouts = report_layouts(deck, pages)
    writers = {'text': add_text, 'table': add_table, 'shape': add_shape,
               'chart': add_chart, 'image': add_image}
    for source in pages:
        mark = next((e for e in source['elements'] if 'footer-mark' in str(e['name'])), None)
        classification = ''.join(r['text'] for line in mark['lines'] for r in line) if mark else ''
        layout = layouts[(source.get('theme', 'report'), classification)]
        slide = deck.slides.add_slide(layout)
        for placeholder in list(slide.placeholders):
            if placeholder != slide.shapes.title:
                slide.shapes._spTree.remove(placeholder._element)
        paint(slide.background.fill, source['background'])
        groups = {}
        for item in source['elements']:
            if item.get('slideNumber') or item['name'] == 'footer-rule':
                continue
            if 'footer-mark' in str(item['name']):
                continue
            target = slide
            if item.get('group') is not None and not item.get('title'):
                key=item['group']
                if key not in groups:
                    groups[key]=slide.shapes.add_group_shape()
                    groups[key].name=item['groupName']
                target=groups[key]
            shape=writers[item['role']](target, item)
            shape.name=(item['groupName']+' · ' if item.get('groupName') else '')+str(item['name'])
        slide.notes_slide.notes_text_frame.text = source['notes']
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('html', type=Path)
    parser.add_argument('pptx', type=Path)
    parser.add_argument('--chrome', type=Path, default=Path('/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'))
    parser.add_argument('--screenshots', type=Path)
    parser.add_argument('--screenshots-only', action='store_true')
    args = parser.parse_args()
    pages = extract(args.html.resolve(strict=True), args.chrome.resolve(strict=True), args.screenshots)
    if not args.screenshots_only:
        convert(pages, args.pptx.resolve())
        print(f'Wrote {len(pages)} editable slides: {args.pptx}')
    if args.screenshots:
        print(f'HTML screenshots: {args.screenshots}')


if __name__ == '__main__':
    main()
